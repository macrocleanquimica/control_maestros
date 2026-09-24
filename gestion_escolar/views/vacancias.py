import os
import openpyxl
from datetime import datetime, date

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.http import JsonResponse, HttpResponse
from django.db import transaction
from django.urls import reverse
from django.conf import settings

from django.db.models import Q
from ..models import LoteReporteVacancia, Vacancia, Prelacion, MotivoTramite, PlantillaTramite, Historial, Interinato
from ..forms import VacanciaForm

from .helpers import get_month_diff, get_full_name, send_to_google_sheet, generate_word_document, serialize_form_data, format_date_for_solicitud_asignacion

@permission_required('gestion_escolar.acceder_vacancias', raise_exception=True)
def gestionar_lote_vacancia(request, lote_id=None):
    if lote_id:
        lote = get_object_or_404(LoteReporteVacancia, id=lote_id, usuario_generador=request.user)
        if lote.estado == 'GENERADO':
            messages.warning(request, "Este lote ya ha sido procesado y no puede ser modificado.")
            return redirect('lista_lotes_vacancia')
        # Cancel any other EN_PROCESO lotes if this one is being explicitly edited
        LoteReporteVacancia.objects.filter(
            usuario_generador=request.user,
            estado='EN_PROCESO'
        ).exclude(id=lote.id).update(estado='CANCELADO')
    else:
        lotes_en_proceso = LoteReporteVacancia.objects.filter(
            usuario_generador=request.user,
            estado='EN_PROCESO'
        ).order_by('-fecha_creacion')

        if lotes_en_proceso.exists():
            lote = lotes_en_proceso.first()
            if lotes_en_proceso.count() > 1:
                for old_lote in lotes_en_proceso[1:]:
                    old_lote.estado = 'CANCELADO'
                    old_lote.save()
        else:
            lote = LoteReporteVacancia.objects.create(
                usuario_generador=request.user,
                estado='EN_PROCESO'
            )

    if request.method == 'POST':
        form = VacanciaForm(request.POST)
        if form.is_valid():
            maestro = form.cleaned_data['maestro_titular']
            escuela = maestro.id_escuela

            vacancia = form.save(commit=False)
            vacancia.lote = lote

            maestro_interino_obj = form.cleaned_data.get('maestro_interino')
            if maestro_interino_obj:
                vacancia.nombre_interino = f'{maestro_interino_obj.nombres} {maestro_interino_obj.a_paterno} {maestro_interino_obj.a_materno}'
                vacancia.curp_interino = maestro_interino_obj.curp
                if maestro_interino_obj.curp:
                    prelacion = Prelacion.objects.filter(curp=maestro_interino_obj.curp).first()
                    if prelacion:
                        vacancia.posicion_orden = prelacion.pos_orden
                        vacancia.folio_prelacion = prelacion.folio

            vacancia.direccion = f"{escuela.nombre_ct}, {escuela.domicilio}, DURANGO, {escuela.region}, {escuela.get_turno_display()}, ZONA ECONOMICA:{escuela.zona_economica}"
            
            apreciacion_desc = form.cleaned_data['apreciacion'].descripcion
            if apreciacion_desc.startswith("ADMISIÓN"):
                vacancia.destino = "Admisión"
            elif apreciacion_desc.startswith("PROMOCIÓN"):
                vacancia.destino = "Promoción vertical"
            else:
                vacancia.destino = ""

            vacancia.sostenimiento = "Federalizado" if escuela.sostenimiento == 'FEDERAL' else "Estatal"
            vacancia.turno = escuela.get_turno_display()
            vacancia.tipo_movimiento_reporte = form.cleaned_data['tipo_movimiento_original']
            
            hrs = maestro.hrs or "00.0"
            vacancia.tipo_plaza = "JORNADA" if hrs == "00.0" else "HORA/SEMANA/MES"
            if hrs == "00.0":
                vacancia.horas = None
            else:
                try:
                    vacancia.horas = str(int(float(hrs)))
                except (ValueError, TypeError):
                    vacancia.horas = None

            vacancia.municipio = escuela.region
            vacancia.zona_economica = f"Zona {escuela.zona_economica}"
            vacancia.categoria = maestro.categog.id_categoria if maestro.categog else ''
            vacancia.clave_presupuestal = maestro.clave_presupuestal
            vacancia.techo_financiero = maestro.techo_f
            vacancia.clave_ct = escuela.id_escuela
            vacancia.nombre_titular_reporte = f'{maestro.nombres} {maestro.a_paterno} {maestro.a_materno}'

            vacancia.save()

            # Registrar automáticamente el interinato en el historial del maestro interino
            if maestro_interino_obj:
                try:
                    Interinato.objects.create(
                        maestro_interino=maestro_interino_obj,
                        maestro_titular=maestro,
                        clave_presupuestal=vacancia.clave_presupuestal,
                        escuela=escuela,
                        techo_financiero=vacancia.techo_financiero,
                        fecha_inicio=vacancia.fecha_inicio,
                        fecha_final=vacancia.fecha_final,
                        motivo=vacancia.tipo_movimiento_original,
                        tipo='INTERINO',
                        estatus='ACTIVO',
                        observaciones=vacancia.observaciones,
                        vacancia=vacancia,
                    )
                except Exception as e:
                    print(f"DEBUG: ❌ Error al registrar interinato en historial: {e}")

            # Se elimina la creación de Historial inmediata por petición del usuario (se hará al exportar USICAMM)
            messages.success(request, "Vacancia agregada al lote actual.")
            if lote_id:
                return redirect('gestionar_lote_vacancia_con_id', lote_id=lote.id)
            else:
                return redirect('gestionar_lote_vacancia')
        else:
            messages.error(request, "Por favor corrige los errores en el formulario.")
    else:
        form = VacanciaForm()

    vacancias_en_lote = Vacancia.objects.filter(lote=lote).order_by('-id')
    context = {
        'form': form,
        'lote': lote,
        'vacancias_en_lote': vacancias_en_lote,
        'titulo': f'Generar Reporte de Vacancia (Lote #{lote.id})'
    }
    return render(request, 'gestion_escolar/gestionar_lote_vacancia.html', context)

def _get_lote_y_vacancias(request, lote_id):
    lote = get_object_or_404(LoteReporteVacancia, id=lote_id, usuario_generador=request.user)
    vacancias = lote.vacancias.all()

    if not vacancias.exists():
        raise ValueError("No hay vacancias en este lote para exportar.")

    if lote.estado == 'GENERADO':
        messages.warning(request, "Este lote ya fue procesado anteriormente.")
        raise ValueError('Este lote ya fue procesado.')

    lote.estado = 'PROCESANDO'
    lote.save()
    return lote, vacancias

@login_required
@transaction.atomic
def exportar_paso_word(request, lote_id):
    try:
        lote, vacancias = _get_lote_y_vacancias(request, lote_id)
    except ValueError as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=400)

    plantilla_solicitud_asignacion = PlantillaTramite.objects.filter(nombre="SOLICITUD DE ASIGNACION").first()
    documentos_word_generados = 0
    word_docs_info = []

    if not plantilla_solicitud_asignacion:
        return JsonResponse({'status': 'warning', 'message': 'Plantilla "SOLICITUD DE ASIGNACION" no encontrada. Saltando paso de Word.', 'word_count': 0, 'word_docs': []})

    for vacancia in vacancias:
        if vacancia.maestro_interino and vacancia.fecha_inicio and vacancia.fecha_final:
            form_data_for_word = {
                'plantilla': plantilla_solicitud_asignacion,
                'maestro_titular': vacancia.maestro_titular,
                'maestro_interino': vacancia.maestro_interino,
                'fecha_efecto1': vacancia.fecha_inicio, 
                'fecha_efecto2': vacancia.fecha_final,
                'fecha_efecto3': vacancia.fecha_inicio, 
                'fecha_efecto4': vacancia.fecha_final,
                'folio': vacancia.folio_prelacion, 
                'observaciones': vacancia.observaciones,
                'no_prel_display': vacancia.posicion_orden, 
                'folio_prel_display': vacancia.folio_prelacion,
            }
            
            motivo_tramite_obj = MotivoTramite.objects.filter(motivo_tramite=vacancia.tipo_movimiento_original).first()
            form_data_for_word['motivo_tramite'] = motivo_tramite_obj
            
            tipo_val_display = ''
            if vacancia.maestro_interino.curp:
                prelacion = Prelacion.objects.filter(curp=vacancia.maestro_interino.curp).first()
                if prelacion:
                    tipo_val_display = prelacion.tipo_val
            form_data_for_word['tipo_val_display'] = tipo_val_display

            plantilla_actual = plantilla_solicitud_asignacion
            if vacancia.plantilla_doc == 'NUMERO_36':
                plantilla_36 = PlantillaTramite.objects.filter(nombre="SOLICITUD DE ASIGNACION 36").first()
                if plantilla_36:
                    plantilla_actual = plantilla_36
                else:
                    # Fallback if the specific template object doesn't exist, try to clone and adapt path
                    # but ideally the user should create it in the database.
                    # For now, let's assume it should exist or we can try to find it.
                    print("DEBUG: ⚠️ Plantilla 'SOLICITUD DE ASIGNACION 36' no encontrada en BD. Usando normal.")

            success, doc_path = generate_word_document(form_data_for_word, plantilla_actual, request.user)
            if success:
                try:
                    historial_word = Historial.objects.create(
                        usuario=request.user,
                        tipo_documento=f"Oficio - {plantilla_solicitud_asignacion.nombre}",
                        maestro=vacancia.maestro_titular,
                        ruta_archivo=doc_path,
                        motivo=motivo_tramite_obj.motivo_tramite if motivo_tramite_obj else '',
                        maestro_secundario_nombre=get_full_name(vacancia.maestro_interino),
                        datos_tramite=serialize_form_data(form_data_for_word)
                    )
                    word_docs_info.append({
                        'id': historial_word.id,
                        'nombre': os.path.basename(doc_path),
                        'url': reverse('descargar_archivo_historial', args=[historial_word.id])
                    })
                    documentos_word_generados += 1
                except Exception as e:
                    print(f"DEBUG: ❌ Error creando historial para Word: {e}")
    
    return JsonResponse({
        'status': 'success',
        'message': f'Se generaron {documentos_word_generados} documento(s) Word.',
        'word_count': documentos_word_generados,
        'word_docs': word_docs_info
    })

@login_required
@transaction.atomic
def exportar_paso_gsheets(request, lote_id):
    try:
        lote, vacancias = _get_lote_y_vacancias(request, lote_id)
    except ValueError as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=400)

    try:
        vacancias_enviadas = 0
        errores_gsheets = []
        for vacancia in vacancias:
            # Datos para Google Sheets - se envían todas las vacancias
            google_sheet_row_data = [
                str(vacancias.filter(id__lte=vacancia.id).count()),
                datetime.now().strftime("%Y-%m-%d"),
                "EDUCACIÓN ESPECIAL",
                "Durango",
                vacancia.municipio or '',
                vacancia.direccion or '',
                vacancia.region or '',
                vacancia.zona_economica or '',
                vacancia.destino or '',
                vacancia.apreciacion.descripcion if vacancia.apreciacion else '',
                vacancia.get_tipo_vacante_display() or '',
                vacancia.tipo_plaza or '',
                vacancia.horas if vacancia.tipo_plaza == "HORA/SEMANA/MES" else '',
                vacancia.sostenimiento or '',
                vacancia.fecha_inicio.strftime("%Y-%m-%d") if vacancia.fecha_inicio else '',
                vacancia.fecha_final.strftime("%Y-%m-%d") if vacancia.fecha_final else '',
                vacancia.categoria or '',
                vacancia.pseudoplaza or '',
                vacancia.clave_presupuestal or '',
                vacancia.techo_financiero or '',
                vacancia.clave_ct or '',
                vacancia.turno or '',
                vacancia.tipo_movimiento_original or '',
                            vacancia.observaciones or '',
                            vacancia.curp_interino or '', # Usar el campo pre-populado en Vacancia, que ya contiene el CURP si se ha asignado un interino, o es vacío en caso contrario.
                            '',
                            (vacancia.maestro_interino.form_academica if vacancia.maestro_interino else '') or '', # Acceso seguro para form_academica si maestro_interino no es None.
                                            "N/A" if vacancia.tipo_plaza == "JORNADA" else (vacancia.apreciacion.descripcion if vacancia.apreciacion else ''),
                '', '', '', '',
                format_date_for_solicitud_asignacion(vacancia.fecha_inicio) if vacancia.fecha_inicio else '',
                format_date_for_solicitud_asignacion(vacancia.fecha_final) if vacancia.fecha_final else '',
                '', '',
                f"DEE/{vacancia.folio_prelacion}/2025" if vacancia.folio_prelacion else '',
                '', '',
            ]
            success_gs, message_gs = send_to_google_sheet(google_sheet_row_data)
            if not success_gs:
                error_msg = f"Fallo al enviar datos del interino '{get_full_name(vacancia.maestro_interino)}': {message_gs}"
                errores_gsheets.append(error_msg)
            else:
                vacancias_enviadas += 1
        
        mensaje_final = f'Se enviaron datos de {vacancias_enviadas} vacancia(s) a Google Sheets.'
        if errores_gsheets:
            mensaje_final += f' Hubo {len(errores_gsheets)} error(es).'

        return JsonResponse({
            'status': 'success',
            'message': mensaje_final,
            'gsheets_count': vacancias_enviadas,
            'gsheets_errors': errores_gsheets
        })

    except Exception as e:
        import traceback
        print(f"DEBUG: ❌ Error en exportar_paso_gsheets: {str(e)}")
        print(f"DEBUG: Traceback completo: {traceback.format_exc()}")
        # Si ocurre un error, revertir el estado del lote para que se pueda reintentar
        lote.estado = 'EN_PROCESO'
        lote.save()
        return JsonResponse({'status': 'error', 'message': f"Error interno del servidor al procesar Google Sheets: {str(e)}"}, status=500)

@login_required
@transaction.atomic
def exportar_paso_excel(request, lote_id):
    try:
        lote, vacancias = _get_lote_y_vacancias(request, lote_id)
        
        template_path = os.path.join(settings.BASE_DIR, 'gestion_escolar', 'templates', 'tramites', 'Plantillas', 'Excel', 'FORMATOVACANCIAUSICAMM.xlsx')
        if not os.path.exists(template_path):
            # Respaldo temporal: antigua carpeta raíz
            template_path = os.path.join(settings.BASE_DIR, 'tramites', 'Plantillas', 'Excel', 'FORMATOVACANCIAUSICAMM.xlsx')
        
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"No se encontró el template en la ruta canónica ni en la legacy: {template_path}")

        workbook = openpyxl.load_workbook(template_path)
        sheet = workbook.active
        
        row_num = 2
        campos_excel = [
            'nivel', 'entidad', 'municipio', 'direccion', 'region', 'zona_economica', 'destino', 'apreciacion',
            'tipo_vacante', 'tipo_plaza', 'horas', 'sostenimiento', 'fecha_inicio', 'fecha_final', 'categoria',
            'pseudoplaza', 'clave_presupuestal', 'techo_financiero', 'clave_ct', 'turno', 'tipo_movimiento_reporte', 'observaciones',
            'posicion_orden', 'folio_prelacion', 'curp_interino', 'nombre_interino'
        ]
        
        for vacancia in vacancias:
            for i, field_name in enumerate(campos_excel):
                cell = sheet.cell(row=row_num, column=i + 1)
                valor = ''
                
                try:
                    if field_name == 'curp_interino': 
                        valor = vacancia.maestro_interino.curp if vacancia.maestro_interino else vacancia.curp_interino or ''
                    elif field_name == 'nombre_interino': 
                        valor = get_full_name(vacancia.maestro_interino) if vacancia.maestro_interino else vacancia.nombre_interino or ''
                    elif field_name == 'apreciacion': 
                        valor = vacancia.apreciacion.descripcion if vacancia.apreciacion else ''
                    elif field_name == 'tipo_vacante': 
                        valor = vacancia.get_tipo_vacante_display().capitalize() if vacancia.tipo_vacante else ''
                    elif field_name == 'zona_economica':
                        valor = getattr(vacancia, field_name, '') or ''
                        if valor == 'Zona II':
                            valor = 'Zona 2'
                        elif valor == 'Zona III':
                            valor = 'Zona 3'
                    else: 
                        valor = getattr(vacancia, field_name, '') or ''
                        
                    if field_name in ['fecha_inicio', 'fecha_final'] and valor and isinstance(valor, (datetime, date)):
                        valor = valor.strftime("%Y-%m-%d")
                            
                except Exception as field_error:
                    valor = f"Error: {field_error}"
                
                cell.value = valor
                
            row_num += 1

        output_dir = os.path.join(settings.MEDIA_ROOT, 'reportes_vacancias')
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_filename = f"VACANCIA_{timestamp}.xlsx"
        output_path_server = os.path.join(output_dir, output_filename)
        
        workbook.save(output_path_server)

        historial_excel = Historial.objects.create(
            usuario=request.user, 
            tipo_documento="Reporte de Vacancia", 
            maestro=None,
            ruta_archivo=output_path_server, 
            motivo="Reporte de Vacancia", 
            lote_reporte=lote
        )
        
        # Crear registros individuales para el Kardex de cada maestro (Oficialización del reporte)
        for v in vacancias:
            if v.maestro_titular:
                datos_v = {
                    'tipo_vacante': v.get_tipo_vacante_display(),
                    'tipo_movimiento': v.tipo_movimiento_reporte or v.tipo_movimiento_original,
                    'fecha_inicio': v.fecha_inicio.strftime('%Y-%m-%d') if v.fecha_inicio else '',
                    'fecha_final': v.fecha_final.strftime('%Y-%m-%d') if v.fecha_final else 'AL COBRO',
                    'techo_financiero': v.techo_financiero or '',
                    'clave_presupuestal': v.clave_presupuestal or '',
                    'centro_trabajo': v.clave_ct or ''
                }
                Historial.objects.create(
                    usuario=request.user,
                    tipo_documento="Reporte de Vacancia (Individual)",
                    maestro=v.maestro_titular,
                    ruta_archivo=output_path_server,
                    motivo="Reporte de Vacancia Oficializado",
                    lote_reporte=lote,
                    datos_tramite=datos_v
                )
        
        lote.archivo_generado = os.path.join('reportes_vacancias', output_filename)
        lote.estado = 'GENERADO'
        lote.fecha_generado = datetime.now()
        lote.save()
        response_data = {
            'status': 'success',
            'message': '¡Lote procesado satisfactoriamente!',
            'excel_url': reverse('descargar_archivo_historial', args=[historial_excel.id]),
            'excel_name': output_filename
        }
        return JsonResponse(response_data)

    except Exception as e:
        import traceback
        print(f"DEBUG: ❌ Error generando Excel: {str(e)}")
        print(f"DEBUG: Traceback completo: {traceback.format_exc()}")
        
        messages.error(request, f"Error al generar el archivo Excel: {str(e)}")
        lote.estado = 'EN_PROCESO'
        lote.save()
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)

@login_required
@transaction.atomic
def exportar_paso_excel_eo(request, lote_id):
    """
    Exportación usando el formato formatovalidacionEO.xlsx (Validación EO).
    Cobra: Escuela asociada a maestro.techo_f
    Labora: Escuela asociada a maestro.id_escuela
    """
    try:
        import io
        import zipfile
        from ..models import Escuela
        lote = get_object_or_404(LoteReporteVacancia, id=lote_id, usuario_generador=request.user)
        vacancias = lote.vacancias.all()
        
        if not vacancias.exists():
            return JsonResponse({'status': 'error', 'message': "No hay vacancias en este lote."}, status=400)
        
        template_path = os.path.join(settings.BASE_DIR, 'gestion_escolar', 'templates', 'tramites', 'Plantillas', 'Excel', 'formatovalidacionEO.xlsx')
        if not os.path.exists(template_path):
            # Respaldo temporal: antigua carpeta raíz (sin rutas absolutas C:\).
            template_path = os.path.join(settings.BASE_DIR, 'tramites', 'Plantillas', 'Excel', 'formatovalidacionEO.xlsx')
            if not os.path.exists(template_path):
                raise FileNotFoundError(f"No se encontró el template EO en la ruta canónica ni en la legacy.")

        # Preparar el buffer para el ZIP y directorios de salida
        output_dir = os.path.join(settings.MEDIA_ROOT, 'reportes_vacancias_eo')
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        zip_buffer = io.BytesIO()
        
        with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
            # Mapeo de Zona Económica
            map_ze = {'I': '1', 'II': '2', 'III': '3', 'IV': '4', 'V': '5'}
            def clean_ze(ze):
                if not ze: return ""
                ze_clean = str(ze).strip().upper()
                return map_ze.get(ze_clean, ze_clean)

            # Fecha de Elaboración en D44
            meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
            hoy = datetime.now()
            fecha_texto = f"{hoy.day} de {meses[hoy.month-1]} del {hoy.year}"

            for idx, vacancia in enumerate(vacancias):
                try:
                    maestro = vacancia.maestro_titular
                    if not maestro: continue
                    
                    # Cargar un workbook "limpio" de la plantilla para cada maestro
                    workbook = openpyxl.load_workbook(template_path)
                    sheet = workbook.active
                    
                    escuela_labora = maestro.id_escuela
                    escuela_cobra = Escuela.objects.filter(id_escuela=maestro.techo_f).first() if maestro.techo_f else escuela_labora
                    if not escuela_cobra: escuela_cobra = escuela_labora

                    def get_esc_val(esc, attr, default=""):
                        if not esc: return default
                        return getattr(esc, attr, default) or default

                    def get_esc_turno(esc):
                        return (esc.get_turno_display() or "").upper() if hasattr(esc, 'get_turno_display') else ""

                    # RichText Bold
                    try:
                        from openpyxl.cell.rich_text import CellRichText, TextBlock
                        from openpyxl.cell.text import InlineFont
                        bold_f = InlineFont(b=True)
                        def rich_val(label, val):
                            return CellRichText([TextBlock(bold_f, label), str(val)])
                    except ImportError:
                        def rich_val(label, val):
                            return f"{label}{val}"

                    # Llenar datos (siempre en filas fijas 8, 9, 15-21, 44)
                    sheet["C8"] = (vacancia.get_tipo_vacante_display() or "").upper()
                    sheet["C9"] = "FEDERAL" if get_esc_val(escuela_labora, 'sostenimiento') == 'FEDERAL' else "ESTATAL"
                    sheet["B16"] = "DURANGO"
                    sheet["D44"] = fecha_texto
                    
                    # COBRA (C15-C21)
                    mun_cobra = get_esc_val(escuela_cobra, 'region')
                    if mun_cobra and "DGO" not in mun_cobra.upper(): mun_cobra += " , DGO."
                    sheet["C15"] = rich_val("C.C.T ", get_esc_val(escuela_cobra, 'id_escuela'))
                    sheet["C16"] = rich_val("NOMBRE DEL C.T.: ", get_esc_val(escuela_cobra, 'nombre_ct'))
                    sheet["C17"] = rich_val("TURNO: ", get_esc_turno(escuela_cobra))
                    sheet["C18"] = rich_val("MUNICIPIO: ", mun_cobra)
                    sheet["C19"] = rich_val("LOCALIDAD: ", get_esc_val(escuela_cobra, 'region'))
                    sheet["C20"] = rich_val("DOMICILIO: ", get_esc_val(escuela_cobra, 'domicilio'))
                    sheet["C21"] = rich_val("ZONA ECONOMICA: ", clean_ze(get_esc_val(escuela_cobra, 'zona_economica')))
                    
                    # LABORA (D15-D21)
                    mun_labora = get_esc_val(escuela_labora, 'region')
                    if mun_labora and "DGO" not in mun_labora.upper(): mun_labora += " , DGO."
                    sheet["D15"] = rich_val("C.C.T. ", get_esc_val(escuela_labora, 'id_escuela'))
                    sheet["D16"] = rich_val("NOMBRE DEL C.T. ", get_esc_val(escuela_labora, 'nombre_ct'))
                    sheet["D17"] = rich_val("TURNO: ", get_esc_turno(escuela_labora))
                    sheet["D18"] = rich_val("MUNICIPIO: ", mun_labora)
                    sheet["D19"] = rich_val("LOCALIDAD: ", get_esc_val(escuela_labora, 'region'))
                    sheet["D20"] = rich_val("DOMICILIO: ", get_esc_val(escuela_labora, 'domicilio'))
                    sheet["D21"] = rich_val("ZONA ECONOMICA: ", clean_ze(get_esc_val(escuela_labora, 'zona_economica')))
                    
                    # Otros
                    sheet["E16"] = (vacancia.destino or "ADMISIÓN").upper()
                    f_init = vacancia.fecha_inicio.strftime('%d/%m/%Y') if vacancia.fecha_inicio else ""
                    f_end = vacancia.fecha_final.strftime('%d/%m/%Y') if vacancia.fecha_final else "AL COBRO"
                    sheet["F16"] = f"{f_init} AL {f_end}"
                    sheet["G16"] = vacancia.clave_presupuestal or ""
                    sheet["H16"] = (vacancia.tipo_movimiento_reporte or vacancia.tipo_movimiento_original or "").upper()
                    sheet["I16"] = (vacancia.observaciones or "").upper()
                    
                    # Guardar excel a buffer para el ZIP
                    excel_buffer = io.BytesIO()
                    workbook.save(excel_buffer)
                    
                    # Guardar excel a disco de forma individual para el Kardex
                    safe_name = "".join([c if c.isalnum() else "_" for c in get_full_name(maestro)])
                    individual_filename = f"EO_{safe_name}_{timestamp}.xlsx"
                    individual_path = os.path.join(output_dir, individual_filename)
                    
                    with open(individual_path, "wb") as f:
                        f.write(excel_buffer.getvalue())

                    # Crear registro individual en Historial para el Kardex
                    Historial.objects.create(
                        usuario=request.user,
                        tipo_documento="Validación EO (Individual)",
                        maestro=maestro,
                        ruta_archivo=individual_path,
                        motivo="Generación de Formato EO",
                        lote_reporte=lote
                    )
                    
                    # Añadir al ZIP
                    zip_file.writestr(f"VALIDACION_EO_{safe_name}.xlsx", excel_buffer.getvalue())
                    
                except Exception as row_e:
                    print(f"DEBUG: ❌ Error procesando maestro {idx}: {str(row_e)}")
                    continue

        # Guardar ZIP consolidado
        output_filename = f"VALIDACION_EO_LOTE_{lote.id}_{timestamp}.zip"
        output_path_server = os.path.join(output_dir, output_filename)
        
        with open(output_path_server, "wb") as f:
            f.write(zip_buffer.getvalue())

        historial_zip = Historial.objects.create(
            usuario=request.user, 
            tipo_documento="Paquete Validación EO (ZIP)", 
            maestro=None,
            ruta_archivo=output_path_server, 
            motivo="Exportación Formato EO Individual", 
            lote_reporte=lote
        )
        
        return JsonResponse({
            'status': 'success',
            'message': f'Se generaron archivos para {vacancias.count()} maestros en un ZIP.',
            'excel_url': reverse('descargar_archivo_historial', args=[historial_zip.id]),
            'excel_name': output_filename
        })

    except Exception as e:
        import traceback
        error_msg = traceback.format_exc()
        print(f"DEBUG: ❌ Error generando ZIP EO: {str(e)}")
        print(f"DEBUG: Traceback completo: {error_msg}")
        
        try:
            if 'lote' in locals() and lote:
                lote.estado = 'EN_PROCESO'
                lote.save()
        except: pass

        return JsonResponse({'status': 'error', 'message': f"Error ZIP EO: {str(e)}"}, status=500)

@login_required
def eliminar_vacancia_lote(request, pk):
    if request.method == 'POST':
        vacancia = get_object_or_404(Vacancia, pk=pk)
        if vacancia.lote.usuario_generador == request.user and vacancia.lote.estado == 'EN_PROCESO':
            vacancia.delete()
            messages.success(request, "Vacancia eliminada del lote correctamente.")
            return JsonResponse({'status': 'success'})
        else:
            return JsonResponse({'status': 'error', 'message': 'No autorizado para eliminar esta vacancia.'}, status=403)
    return JsonResponse({'status': 'error', 'message': 'Método no permitido.'}, status=405)

@login_required
@permission_required('gestion_escolar.acceder_reportes', raise_exception=True)
def reporte_vacancias_fecha(request):
    """
    Vista para el reporte detallado de vacancias con filtros por fecha.
    Incluye búsqueda de trámites asociados en el historial.
    """
    f_inicio_desde = request.GET.get('f_inicio_desde')
    f_inicio_hasta = request.GET.get('f_inicio_hasta')
    f_final_desde = request.GET.get('f_final_desde')
    f_final_hasta = request.GET.get('f_final_hasta')
    f_reporte_desde = request.GET.get('f_reporte_desde')
    f_reporte_hasta = request.GET.get('f_reporte_hasta')

    vacancias = Vacancia.objects.select_related('maestro_titular', 'maestro_interino', 'lote', 'maestro_titular__id_escuela').all()

    if f_inicio_desde:
        vacancias = vacancias.filter(fecha_inicio__gte=f_inicio_desde)
    if f_inicio_hasta:
        vacancias = vacancias.filter(fecha_inicio__lte=f_inicio_hasta)
    
    if f_final_desde:
        vacancias = vacancias.filter(fecha_final__gte=f_final_desde)
    if f_final_hasta:
        vacancias = vacancias.filter(fecha_final__lte=f_final_hasta)

    if f_reporte_desde:
        vacancias = vacancias.filter(lote__fecha_creacion__date__gte=f_reporte_desde)
    if f_reporte_hasta:
        vacancias = vacancias.filter(lote__fecha_creacion__date__lte=f_reporte_hasta)

    vacancias = vacancias.order_by('-lote__fecha_creacion', '-id')

    # Lógica para buscar trámites asociados
    for v in vacancias:
        # Buscar en historial trámites del titular posteriores a la fecha del reporte
        potential_tramites = Historial.objects.filter(
            maestro=v.maestro_titular,
            fecha_creacion__gte=v.lote.fecha_creacion,
            tipo_documento__icontains='Trámite'
        )
        
        match = None
        v_cp = v.clave_presupuestal.strip() if v.clave_presupuestal else ""
        for t in potential_tramites:
            data = t.datos_tramite or {}
            # Revisar ambas posibles llaves del JSON (Oficios vs Trámites)
            cp1 = str(data.get('CLAVE_PRESUPUESTAL_TITULAR') or "").strip()
            cp2 = str(data.get('clave_presupuestal_titular_display') or "").strip()
            
            if v_cp and (v_cp == cp1 or v_cp == cp2):
                match = t
                break
        
        v.tramite_vinculado = match

    context = {
        'vacancias': vacancias,
        'titulo': 'Reporte Detallado de Vacancias',
        'filtros': {
            'f_inicio_desde': f_inicio_desde,
            'f_inicio_hasta': f_inicio_hasta,
            'f_final_desde': f_final_desde,
            'f_final_hasta': f_final_hasta,
            'f_reporte_desde': f_reporte_desde,
            'f_reporte_hasta': f_reporte_hasta,
        }
    }
    return render(request, 'gestion_escolar/reporte_vacancias_fecha.html', context)

@login_required
@permission_required('gestion_escolar.acceder_reportes', raise_exception=True)
def exportar_reporte_vacancias_excel(request):
    """
    Genera un archivo Excel con el reporte filtrado de vacancias, incluyendo info de trámites.
    """
    f_inicio_desde = request.GET.get('f_inicio_desde')
    f_inicio_hasta = request.GET.get('f_inicio_hasta')
    f_final_desde = request.GET.get('f_final_desde')
    f_final_hasta = request.GET.get('f_final_hasta')
    f_reporte_desde = request.GET.get('f_reporte_desde')
    f_reporte_hasta = request.GET.get('f_reporte_hasta')

    vacancias = Vacancia.objects.select_related('maestro_titular', 'maestro_interino', 'lote', 'maestro_titular__id_escuela').all()

    if f_inicio_desde:
        vacancias = vacancias.filter(fecha_inicio__gte=f_inicio_desde)
    if f_inicio_hasta:
        vacancias = vacancias.filter(fecha_inicio__lte=f_inicio_hasta)
    
    if f_final_desde:
        vacancias = vacancias.filter(fecha_final__gte=f_final_desde)
    if f_final_hasta:
        vacancias = vacancias.filter(fecha_final__lte=f_final_hasta)

    if f_reporte_desde:
        vacancias = vacancias.filter(lote__fecha_creacion__date__gte=f_reporte_desde)
    if f_reporte_hasta:
        vacancias = vacancias.filter(lote__fecha_creacion__date__lte=f_reporte_hasta)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte Vacancias"

    headers = [
        "CLAVE PRESUPUESTAL", "FECHA REPORTE", "FECHA INICIO", "FECHA FINAL", 
        "TITULAR", "INTERINO", "CURP INTERINO", "TIPO VACANTE", 
        "MOTIVO", "CCT", "ESCUELA", "OBSERVACIONES", "TRÁMITE REALIZADO", "FECHA TRÁMITE"
    ]
    ws.append(headers)

    for v in vacancias.order_by('-lote__fecha_creacion', '-id'):
        # Buscar trámite vinculado de forma robusta
        potential_tramites = Historial.objects.filter(
            maestro=v.maestro_titular,
            fecha_creacion__gte=v.lote.fecha_creacion,
            tipo_documento__icontains='Trámite'
        )
        
        tramite = None
        v_cp = v.clave_presupuestal.strip() if v.clave_presupuestal else ""
        for t in potential_tramites:
            data = t.datos_tramite or {}
            cp1 = str(data.get('CLAVE_PRESUPUESTAL_TITULAR') or "").strip()
            cp2 = str(data.get('clave_presupuestal_titular_display') or "").strip()
            if v_cp and (v_cp == cp1 or v_cp == cp2):
                tramite = t
                break

        nombre_interino = ""
        if v.maestro_interino:
            nombre_interino = f"{v.maestro_interino.nombres} {v.maestro_interino.a_paterno} {v.maestro_interino.a_materno}"
        else:
            nombre_interino = v.nombre_interino or ""

        row = [
            v.clave_presupuestal or "",
            v.lote.fecha_creacion.strftime("%d/%m/%Y") if v.lote and v.lote.fecha_creacion else "",
            v.fecha_inicio.strftime("%d/%m/%Y") if v.fecha_inicio else "",
            v.fecha_final.strftime("%d/%m/%Y") if v.fecha_final else "AL COBRO",
            v.nombre_titular_reporte or "",
            nombre_interino,
            v.curp_interino or "",
            v.get_tipo_vacante_display() or "",
            v.tipo_movimiento_reporte or v.tipo_movimiento_original or "",
            v.clave_ct or "",
            v.maestro_titular.id_escuela.nombre_ct if v.maestro_titular and v.maestro_titular.id_escuela else "",
            v.observaciones or "",
            tramite.tipo_documento if tramite else "NO REALIZADO",
            tramite.fecha_creacion.strftime("%d/%m/%Y %H:%M") if tramite else "-"
        ]
        ws.append(row)

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 50)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename="reporte_vacancias_completo.xlsx"'},
    )
    wb.save(response)
    return response

@login_required
@permission_required('gestion_escolar.acceder_vacancias', raise_exception=True)
def lista_lotes_vacancia(request):
    lotes = LoteReporteVacancia.objects.filter(usuario_generador=request.user).order_by('-fecha_creacion')
    context = {
        'lotes': lotes,
        'titulo': 'Histórico de Lotes de Vacancia'
    }
    return render(request, 'gestion_escolar/lista_lotes_vacancia.html', context)

@login_required
@permission_required('gestion_escolar.acceder_vacancias', raise_exception=True)
@transaction.atomic
def cancelar_lote_vacancia(request, lote_id):
    if request.method == 'POST':
        try:
            lote = get_object_or_404(LoteReporteVacancia, id=lote_id, usuario_generador=request.user)
            if lote.estado == 'EN_PROCESO':
                lote.estado = 'CANCELADO'
                lote.save()
                messages.success(request, f"Lote #{lote.id} cancelado exitosamente.")
                return JsonResponse({'status': 'success', 'message': f"Lote #{lote.id} cancelado exitosamente."})
            else:
                return JsonResponse({'status': 'error', 'message': f"El lote #{lote.id} no puede ser cancelado en su estado actual ({lote.estado})."}, status=400)
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': f"Error al cancelar el lote: {str(e)}"}, status=500)
    return JsonResponse({'status': 'error', 'message': 'Método no permitido.'}, status=405)