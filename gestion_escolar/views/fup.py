from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.urls import reverse
from django.db.models import Q
from django.http import JsonResponse, HttpResponse
from ..models import FUP, Maestro, Escuela, Zona
from ..forms import FUPForm
import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import date

@login_required
def lista_fup(request):
    return render(request, 'gestion_escolar/lista_fup.html')

@login_required
def fup_datatable_ajax(request):
    draw = int(request.GET.get('draw', 0))
    start = int(request.GET.get('start', 0))
    length = int(request.GET.get('length', 10))
    search_value = request.GET.get('search[value]', '')

    order_column_index = int(request.GET.get('order[0][column]', 0))
    order_dir = request.GET.get('order[0][dir]', 'asc')
    column_names = ['folio', 'fecha', 'nombre_completo', 'rfc', 'clave_presupuestal', 'techo_financiero', 'efectos']
    order_column = column_names[order_column_index]
    if order_dir == 'desc':
        order_column = f'-{order_column}'

    queryset = FUP.objects.select_related('maestro').all()
    records_total = queryset.count()

    if search_value:
        from unidecode import unidecode
        search_unaccented = unidecode(search_value.upper()).strip()
        
        # Dividir el texto de búsqueda en palabras
        search_words = search_unaccented.split()
        
        # Crear query base para campos simples
        query = Q(folio__icontains=search_value) | \
                Q(rfc__icontains=search_value) | \
                Q(clave_presupuestal__icontains=search_value) | \
                Q(techo_financiero__icontains=search_value)
        
        # Para el nombre completo, buscar que todas las palabras estén presentes
        if search_words:
            nombre_query = Q()
            for word in search_words:
                nombre_query &= Q(nombre_completo__icontains=word)
            query |= nombre_query
        
        queryset = queryset.filter(query)

    records_filtered = queryset.count()
    queryset = queryset.order_by(order_column)[start:start + length]

    data = []
    for fup in queryset:
        
        pdf_button = ''
        if fup.archivo:
            pdf_button = f'<a href="{fup.archivo.url}" class="btn btn-light border btn-action-custom text-danger" target="_blank" title="Ver PDF"><i class="fas fa-file-pdf"></i></a>'

        actions = '<div class="btn-group" role="group">'
        actions += f'<a href="{reverse("detalle_fup", args=[fup.pk])}" class="btn btn-light border btn-action-custom text-info" title="Detalle"><i class="fas fa-eye"></i></a>'
        
        if request.user.has_perm('gestion_escolar.change_fup'):
            actions += f'<a href="{reverse("editar_fup", args=[fup.pk])}" class="btn btn-light border btn-action-custom text-primary" title="Editar"><i class="fas fa-edit"></i></a>'

        if request.user.has_perm('gestion_escolar.delete_fup'):
            actions += f'<a href="{reverse("eliminar_fup", args=[fup.pk])}" class="btn btn-light border btn-action-custom text-danger" title="Eliminar"><i class="fas fa-trash"></i></a>'
        
        actions += '</div>'
        
        data.append([
            fup.folio,
            fup.fecha.strftime('%Y-%m-%d'),
            fup.nombre_completo,
            fup.rfc,
            fup.clave_presupuestal,
            fup.techo_financiero,
            fup.efectos,
            pdf_button,
            actions
        ])

    response = {
        'draw': draw,
        'recordsTotal': records_total,
        'recordsFiltered': records_filtered,
        'data': data,
    }
    return JsonResponse(response)

@login_required
def crear_fup(request):
    if request.method == 'POST':
        form = FUPForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            messages.success(request, 'FUP creado correctamente.')
            return redirect('lista_fup')
        else:
            messages.error(request, 'Por favor corrige los errores.')
    else:
        form = FUPForm()
    return render(request, 'gestion_escolar/form_fup.html', {'form': form, 'titulo': 'Captura de FUP'})

@login_required
def editar_fup(request, pk):
    fup = get_object_or_404(FUP, pk=pk)
    if request.method == 'POST':
        form = FUPForm(request.POST, request.FILES, instance=fup)
        if form.is_valid():
            form.save()
            messages.success(request, 'FUP actualizado correctamente.')
            return redirect('lista_fup')
        else:
            messages.error(request, 'Por favor corrige los errores.')
    else:
        form = FUPForm(instance=fup)
    return render(request, 'gestion_escolar/form_fup.html', {'form': form, 'titulo': 'Editar FUP'})

@login_required
def eliminar_fup(request, pk):
    fup = get_object_or_404(FUP, pk=pk)
    if request.method == 'POST':
        fup.delete()
        messages.success(request, 'FUP eliminado correctamente.')
        return redirect('lista_fup')
    return render(request, 'gestion_escolar/eliminar_fup.html', {'fup': fup})

@login_required
def detalle_fup(request, pk):
    fup = get_object_or_404(FUP, pk=pk)
    
    zona_techo_financiero = None
    if fup.techo_financiero:
        try:
            escuela = Escuela.objects.get(id_escuela=fup.techo_financiero)
            zona_techo_financiero = escuela.zona_esc.numero
        except Escuela.DoesNotExist:
            pass
            
    return render(request, 'gestion_escolar/detalle_fup.html', {
        'fup': fup, 
        'zona_techo_financiero': zona_techo_financiero
    })

@login_required
def get_maestro_data_fup(request):
    maestro_id = request.GET.get('maestro_id')
    if not maestro_id:
        return JsonResponse({'error': 'No se proporcionó ID del maestro'}, status=400)
    
    try:
        maestro = Maestro.objects.get(pk=maestro_id)
        data = {
            'rfc': maestro.rfc,
            'curp': maestro.curp,
            'clave_presupuestal': maestro.clave_presupuestal,
            'techo_financiero': maestro.techo_f,
        }
        return JsonResponse(data)
    except Maestro.DoesNotExist:
        return JsonResponse({'error': 'Maestro no encontrado'}, status=404)
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@login_required
def exportar_fup_excel(request):
    """Exporta la lista de FUPs a un archivo Excel."""
    filtro = request.GET.get('filtro', '')
    fup_qs = FUP.objects.select_related('maestro').all().order_by('-fecha')

    if filtro:
        from unidecode import unidecode
        search_unaccented = unidecode(filtro.upper()).strip()
        search_words = search_unaccented.split()
        
        query = Q(folio__icontains=filtro) | \
                Q(rfc__icontains=filtro) | \
                Q(clave_presupuestal__icontains=filtro) | \
                Q(techo_financiero__icontains=filtro)
        
        if search_words:
            nombre_query = Q()
            for word in search_words:
                nombre_query &= Q(nombre_completo__icontains=word)
            query |= nombre_query
        
        fup_qs = fup_qs.filter(query)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte FUPs"

    # Pre-cargar mapa de zonas escolares: {CCT (id_escuela): numero_zona}
    zona_map = {
        e.id_escuela: e.zona_esc.numero 
        for e in Escuela.objects.select_related('zona_esc').all() 
        if e.zona_esc
    }


    headers = [
        "Folio", "Fecha", "Nombre Completo", "RFC", "Clave Presupuestal", 
        "Techo Financiero", "Zona Escolar", "Efectos", "Sostenimiento", "Observaciones"
    ]
    ws.append(headers)
    
    # Apply styles to headers
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for fup in fup_qs:
        row = [
            fup.folio,
            fup.fecha.strftime("%Y-%m-%d") if isinstance(fup.fecha, date) else fup.fecha,
            fup.nombre_completo,
            fup.rfc,
            fup.clave_presupuestal,
            fup.techo_financiero,
            zona_map.get(fup.techo_financiero, ''), # Zona Escolar
            fup.efectos,
            fup.sostenimiento,
            fup.observaciones,
        ]
        ws.append(row)

    # Adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 15 # Zona Escolar
    ws.column_dimensions['H'].width = 30 # Efectos
    ws.column_dimensions['I'].width = 15 # Sostenimiento
    ws.column_dimensions['J'].width = 50 # Observaciones

    # Generar nombre de archivo con fecha actual
    from datetime import datetime
    fecha_actual = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'reporte_fups_{fecha_actual}.xlsx'
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    wb.save(response)
    return response

@login_required
def reporte_fups_opciones(request):
    """Vista para mostrar el formulario de selección de Zona y Fechas para reporte de FUPs."""
    zonas = Zona.objects.all().order_by('numero')
    
    context = {
        'titulo': 'Reporte de FUPs por Zona',
        'zonas': zonas,
    }
    return render(request, 'gestion_escolar/reportes/reporte_fups_opciones.html', context)


@login_required
def exportar_fups_zona_excel(request):
    """Genera el reporte Excel de FUPs filtrado por Zona y opcionalmente por rango de fechas."""
    zona_id = request.GET.get('zona')
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')

    if not zona_id:
        messages.error(request, 'Debe seleccionar una Zona Escolar.')
        return redirect('reporte_fups_opciones')

    try:
        zona = Zona.objects.get(id=zona_id)
    except (ValueError, Zona.DoesNotExist):
        messages.error(request, 'Zona inválida.')
        return redirect('reporte_fups_opciones')

    # Obtener todas las claves CCT de las escuelas que pertenecen a la zona seleccionada
    ccts_zona = Escuela.objects.filter(zona_esc=zona).values_list('id_escuela', flat=True)

    # Filtrar FUPs donde su techo_financiero sea alguna de las escuelas de esa zona
    fups = FUP.objects.filter(techo_financiero__in=ccts_zona)
    
    # Aplicar filtros de fecha si existen
    if fecha_inicio:
        fups = fups.filter(fecha__gte=fecha_inicio)
    if fecha_fin:
        fups = fups.filter(fecha__lte=fecha_fin)
        
    fups = fups.order_by('fecha', 'id')

    # Configuración del documento Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"FUPs Zona {zona.numero}"

    # Título principal del documento
    titulo_principal = f"REPORTE DE FUPS - ZONA {zona.numero}"
    if fecha_inicio and fecha_fin:
        titulo_principal = f"DEL {fecha_inicio} AL {fecha_fin} ZONA {zona.numero}"
    elif fecha_inicio:
        titulo_principal = f"A PARTIR DE {fecha_inicio} ZONA {zona.numero}"
    elif fecha_fin:
        titulo_principal = f"HASTA EL {fecha_fin} ZONA {zona.numero}"
    else:
        titulo_principal = f"TODOS LOS REGISTROS ZONA {zona.numero}"

    ws.merge_cells('A1:F1')
    celda_titulo = ws['A1']
    celda_titulo.value = titulo_principal
    celda_titulo.font = Font(bold=True, size=14)
    celda_titulo.alignment = Alignment(horizontal='center', vertical='center')
    
    # Fila vacía para separación
    ws.append([])
    
    # Encabezados de tabla
    headers = [
        "Consecutivo", "Nombre", "Folio", "Techo financiero", 
        "Efectos del movimiento", "Observaciones"
    ]
    ws.append(headers)
    
    # Estilo para los encabezados (fila 3 ahora)
    for cell in ws[3]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Rellenar los datos
    consecutivo = 1
    for fup in fups:
        row = [
            consecutivo,
            fup.nombre_completo,
            fup.folio,
            fup.techo_financiero,
            fup.efectos,
            fup.observaciones
        ]
        ws.append(row)
        consecutivo += 1

    # Ajustar ancho de columnas
    ws.column_dimensions['A'].width = 12  # Consecutivo
    ws.column_dimensions['B'].width = 45  # Nombre
    ws.column_dimensions['C'].width = 15  # Folio
    ws.column_dimensions['D'].width = 20  # Techo financiero
    ws.column_dimensions['E'].width = 30  # Efectos
    ws.column_dimensions['F'].width = 40  # Observaciones

    # Preparar la respuesta HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    from datetime import datetime
    fecha_hoy = datetime.now().strftime('%Y%m%d%H%M')
    filename = f'reporte_fups_zona_{zona.numero}_{fecha_hoy}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    wb.save(response)
    return response